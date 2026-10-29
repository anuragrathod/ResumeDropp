"""Unit tests for src/excel_exporter.py."""
import io
import openpyxl
from schemas.candidate import CandidateRecord, EducationRecord, Project, WorkExperience
from src.excel_exporter import export_to_bytes, get_highest_education


def test_highest_education_ranking():
    edu1 = EducationRecord(qualification="B.Tech", graduation_year="2018")
    edu2 = EducationRecord(qualification="M.Tech", graduation_year="2020")
    edu3 = EducationRecord(qualification="Ph.D.", graduation_year="2024")

    highest = get_highest_education([edu1, edu2, edu3])
    assert highest is not None
    assert highest.qualification == "Ph.D."


def test_excel_export_creates_5_sheets():
    cand = CandidateRecord(
        resume_filename="alice_resume.pdf",
        full_name="Alice Johnson",
        email="alice@example.com",
        phone="+1 (650) 555-0101",
        location="Palo Alto, CA",
        education=[EducationRecord(qualification="M.Sc. CS", institution="Stanford", graduation_year="2018")],
        work_experience=[WorkExperience(employer="Acme Inc", job_title="Tech Lead", start_date="2020", end_date="Present")],
        projects=[Project(name="OpenTrace", description="Tracing tool", technologies=["Python", "Kafka"])],
        technical_skills=["Python", "Docker"],
    )

    failed = [{"filename": "corrupt.pdf", "doc_status": "corrupt", "error": "Bad PDF header"}]

    excel_bytes = export_to_bytes([cand], failed_docs=failed)
    assert len(excel_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert sheet_names == ["Candidates", "Education", "Work Experience", "Projects", "Processing Log"]

    # Verify Candidates row
    ws_cand = wb["Candidates"]
    assert ws_cand.max_row == 2
    assert ws_cand.cell(row=2, column=3).value == "Alice Johnson"

    # Verify Log row
    ws_log = wb["Processing Log"]
    assert ws_log.max_row == 3  # 1 header + 1 success + 1 failed
