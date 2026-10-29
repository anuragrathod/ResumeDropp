"""Utility script to create synthetic sample resumes in sample_resumes/."""
from pathlib import Path
import fitz
import docx

sample_dir = Path(__file__).resolve().parent.parent / "sample_resumes"
sample_dir.mkdir(parents=True, exist_ok=True)

# 1. Create alice_chen_resume.pdf
pdf_path = sample_dir / "alice_chen_resume.pdf"
doc = fitz.open()
page = doc.new_page()
alice_text = """Alice Chen
alice.chen@example.com | +1 (415) 555-0142 | San Francisco, CA
https://linkedin.com/in/alicechen | https://github.com/alicechen

SUMMARY
Senior Data Engineer with 7 years experience in distributed pipelines and ETL.

EDUCATION
M.Sc. Data Science — UC Berkeley, 2017
B.Sc. Computer Engineering — San Jose State University, 2015

SKILLS & TOOLS
Python, SQL, Apache Spark, Kafka, AWS, Docker, Kubernetes, PostgreSQL

EXPERIENCE
Lead Data Engineer — DataStream Inc. (Jan 2021 – Present)
  • Architected real-time streaming ETL platform processing 50M events daily.
  • Reduced pipeline compute costs by 35% using Apache Spark optimizations.

Data Engineer — CloudMetrics (Jun 2017 – Dec 2020)
  • Built automated data quality validation pipelines using Python and SQL.

PROJECTS
StreamPulse — Open-source telemetry metrics streamer (Python, Kafka)
"""
page.insert_text((50, 72), alice_text, fontsize=11)
doc.save(str(pdf_path))
doc.close()

# 2. Create bob_patel_resume.docx
docx_path = sample_dir / "bob_patel_resume.docx"
d = docx.Document()
d.add_heading("Bob Patel", level=1)
d.add_paragraph("bob.patel@techcorp.in | +91-98234-56789 | Pune, India")
d.add_paragraph("https://linkedin.com/in/bobpatel | https://github.com/bobpatel")

d.add_heading("Education", level=2)
d.add_paragraph("B.E. Computer Technology — COEP Pune, 2020")

d.add_heading("Technical Skills", level=2)
d.add_paragraph("Java, Spring Boot, Microservices, MySQL, Redis, Docker, Git")

d.add_heading("Experience", level=2)
d.add_paragraph("Backend Developer — TechCorp Solutions (Aug 2020 – Present)")
d.add_paragraph("• Designed and deployed high-throughput microservices using Spring Boot.")
d.add_paragraph("• Integrated Redis caching to reduce database query latency by 50%.")

d.save(str(docx_path))

print(f"Created sample resumes in {sample_dir}")
